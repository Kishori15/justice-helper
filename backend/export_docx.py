"""
DOCX Export Module for JusticeHelper cases.
Implements ARCHITECTURE.md §2.8 and DATA_SCHEMA.md §7.
"""
from io import BytesIO
import docx
from docx.shared import Pt, RGBColor, Inches
from backend.models import CaseObject


def generate_docx_bytes(case: CaseObject) -> bytes:
    """
    Renders case summary and complaint drafts into a formatted DOCX document.
    """
    doc = docx.Document()

    # Title
    title = doc.add_heading("JusticeHelper — Case Summary & Drafts", level=1)
    title.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    # Metadata
    p_meta = doc.add_paragraph()
    p_meta.add_run(f"Case ID: {case.case_id}  |  Date: {case.created_at[:10]}  |  Status: {case.status}\n").bold = True

    # 1. Order & Issue Details
    h1 = doc.add_heading("1. Order & Issue Details", level=2)
    h1.runs[0].font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    p_order = doc.add_paragraph()
    if case.order_info:
        p_order.add_run(f"Platform: {case.order_info.platform} | Product: {case.order_info.product_name}\n")
        p_order.add_run(f"Price Paid: ₹{case.order_info.price_paid} | Payment Mode: {case.order_info.payment_mode}\n")
    if case.issue:
        p_order.add_run(f"Issue Type: {case.issue.issue_type.replace('_', ' ').title()}\n")
        p_order.add_run(f"Description: {case.issue.description}\n")

    # 2. Rights & Legal Grounds
    h2 = doc.add_heading("2. Legal Rights & Statutory Grounds", level=2)
    h2.runs[0].font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    if case.generation_output:
        doc.add_paragraph(case.generation_output.rights_summary)
        doc.add_heading("Grounded Legal Basis:", level=3)
        for item in case.generation_output.legal_basis:
            p_basis = doc.add_paragraph(style="List Bullet")
            p_basis.add_run(f"{item.claim}\n").bold = True
            run_cite = p_basis.add_run(f"Citation: {item.citation} (Ref: {item.snippet_id})")
            run_cite.italic = True
            run_cite.font.color.rgb = RGBColor(0x06, 0x5F, 0x46)

    # 3. Seller Notice Draft
    h3 = doc.add_heading("3. Draft: Seller / Platform Grievance Notice", level=2)
    h3.runs[0].font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    if case.generation_output and case.generation_output.draft_email:
        em = case.generation_output.draft_email
        p_em_subj = doc.add_paragraph()
        p_em_subj.add_run(f"Subject: {em.subject}\n").bold = True
        doc.add_paragraph(em.body)

    # 4. NCH Complaint Draft
    h4 = doc.add_heading("4. Draft: NCH / Consumer Commission Complaint", level=2)
    h4.runs[0].font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    if case.generation_output and case.generation_output.draft_nch_complaint:
        nch = case.generation_output.draft_nch_complaint
        doc.add_paragraph(f"Complainant:\n{nch.complainant_details}\n")
        doc.add_paragraph(f"Opposite Party:\n{nch.opposite_party_details}\n")
        doc.add_paragraph(f"Jurisdiction Note:\n{nch.jurisdiction_note}\n")
        doc.add_paragraph("Statement of Facts:").bold = True
        doc.add_paragraph(nch.facts)
        doc.add_paragraph("Grounds of Complaint:").bold = True
        doc.add_paragraph(nch.grounds)
        doc.add_paragraph("Relief Sought:").bold = True
        doc.add_paragraph(nch.relief_sought)
        if nch.enclosures:
            doc.add_paragraph("List of Enclosures:").bold = True
            for enc in nch.enclosures:
                doc.add_paragraph(enc, style="List Bullet")
        doc.add_paragraph("Verification:").bold = True
        doc.add_paragraph(nch.verification_clause)

    # Disclaimer
    doc.add_paragraph()
    p_disc = doc.add_paragraph()
    disclaimer_text = case.generation_output.disclaimer if case.generation_output else "Informational reference draft only. Not formal legal advice."
    run_disc = p_disc.add_run(f"DISCLAIMER: {disclaimer_text}")
    run_disc.font.size = Pt(8.5)
    run_disc.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
